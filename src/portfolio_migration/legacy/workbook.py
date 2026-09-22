"""Build the legacy "Monthly Portfolio Pack" workbook exactly as Finance runs it today.

The workbook is formula driven: the data tabs hold the pasted DWH extracts and
every KPI is a live Excel formula over them. Nothing here computes a KPI in
Python. The published figures come from recalculating the formulas
(see ``evaluate.py``), so whatever the formulas do, right or wrong, is what the
business has been reading.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from portfolio_migration import config as cfg
from portfolio_migration.generate.core import CoreData
from portfolio_migration.legacy import extracts

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
FIRST_ROW = 5
FMT_MONEY = "#,##0"
FMT_RATE = "0.00%"
FMT_FX = "0.0000"

# Columns of each country tab: (header, number format).
COUNTRY_COLUMNS = [
    ("Month end", "yyyy-mm-dd"),
    ("Portfolio balance (local)", FMT_MONEY),
    ("FX rate to EUR", FMT_FX),
    ("Portfolio balance (EUR)", FMT_MONEY),
    ("New accounts", "#,##0"),
    ("New originations (local)", FMT_MONEY),
    ("New originations (EUR)", FMT_MONEY),
    ("Balance 30+ DPD (local)", FMT_MONEY),
    ("30+ DPD rate", FMT_RATE),
    ("Balance 90+ DPD (local)", FMT_MONEY),
    ("90+ DPD rate", FMT_RATE),
    ("Applications received", "#,##0"),
    ("Applications approved", "#,##0"),
    ("Approval rate", FMT_RATE),
    ("Active accounts", "#,##0"),
    ("Avg balance per customer (EUR)", FMT_MONEY),
]

# The KPI cells that make up the published pack, by country tab column.
PUBLISHED_COUNTRY_KPIS = {
    "D": "portfolio_balance_eur",
    "E": "new_accounts",
    "G": "new_originations_eur",
    "I": "dpd30_rate",
    "K": "dpd90_rate",
    "N": "approval_rate",
    "P": "avg_balance_per_customer_eur",
}
PUBLISHED_SUMMARY_KPIS = {
    "B": "portfolio_balance_eur",
    "C": "new_accounts",
    "D": "new_originations_eur",
    "E": "dpd30_rate",
    "F": "dpd90_rate",
    "G": "approval_rate",
    "H": "avg_balance_per_customer_eur",
}

DB = "Data_Balances"
DO = "Data_Originations"
DA = "Data_Applications"


def _header(ws, row: int, headers: list[str]) -> None:
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")


def _write_frame(ws, df: pd.DataFrame, start_row: int = 1) -> None:
    _header(ws, start_row, list(df.columns))
    for r, rec in enumerate(df.itertuples(index=False), start=start_row + 1):
        for c, v in enumerate(rec, start=1):
            ws.cell(row=r, column=c, value=v.item() if hasattr(v, "item") else v)
    for i in range(1, len(df.columns) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def _readme(ws, as_of: date) -> None:
    ws.column_dimensions["A"].width = 110
    lines = [
        ("Monthly Portfolio Pack", Font(bold=True, size=16)),
        ("Owner: Group Finance, Reporting Team", None),
        (f"Data as of: {as_of:%d %B %Y}", None),
        ("", None),
        ("Monthly process", Font(bold=True)),
        ("1. Receive the three DWH extracts by email (balance cube, originations, applications).", None),
        ("2. Paste each extract below the last row of its Data_ tab. Copy the Product line formula down on Data_Balances.", None),
        ("3. Add the new month end rates from the Treasury email to the FX table on Lookups.", None),
        ("4. On each country tab, copy the last row down one row and enter the new month end date in column A.", None),
        ("5. Extend the Summary tab by one row. Check the chart picks it up.", None),
        ("6. Paste Summary as values into the board pack slides.", None),
        ("", None),
        ("Tabs", Font(bold=True)),
        ("Summary: group KPIs for the board pack.", None),
        ("PL, CZ, RO: country KPIs. Each country analyst maintains their own tab.", None),
        ("Data_Balances, Data_Originations, Data_Applications: pasted DWH extracts. Do not sort.", None),
        ("Lookups: product mapping and Treasury FX rates.", None),
        ("", None),
        ("Change log", Font(bold=True)),
        ("v11  Added Gold card to product mapping.", None),
        ("v12  Split 90+ DPD out of 30+ DPD on country tabs.", None),
        ("v13  Added Card Debt Consolidation Loan to product mapping.", None),
        ("v14  Summary chart re-pointed after row insert.", None),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        if font:
            c.font = font


def _lookups(ws, products: pd.DataFrame, fx: pd.DataFrame) -> tuple[str, str]:
    ws["A1"] = "Product mapping"
    ws["A1"].font = Font(bold=True)
    _header(ws, 2, ["Product code", "Product name", "Product line"])
    for r, p in enumerate(products.itertuples(index=False), start=3):
        ws.cell(row=r, column=1, value=p.product_code)
        ws.cell(row=r, column=2, value=p.product_name)
        ws.cell(row=r, column=3, value=p.product_line)
    prod_last = 2 + len(products)

    ws["E1"] = "Treasury FX, EUR per 1 unit of local currency, month end"
    ws["E1"].font = Font(bold=True)
    for j, h in enumerate(["Month end"] + [c.currency for c in cfg.COUNTRIES], start=5):
        cell = ws.cell(row=2, column=j, value=h)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for r, rec in enumerate(fx.itertuples(index=False), start=3):
        ws.cell(row=r, column=5, value=rec.month_end).number_format = "yyyy-mm-dd"
        for j, cur in enumerate([c.currency for c in cfg.COUNTRIES]):
            cell = ws.cell(row=r, column=6 + j, value=float(getattr(rec, cur)))
            cell.number_format, cell.fill = FMT_FX, INPUT_FILL
    fx_last = 2 + len(fx)
    for col, w in zip("ABCDEFGH", (14, 32, 16, 3, 12, 10, 10, 10)):
        ws.column_dimensions[col].width = w
    last_fx_col = get_column_letter(5 + len(cfg.COUNTRIES))
    fx_formula = (f"=INDEX(Lookups!$F$3:${last_fx_col}${fx_last},MATCH($A{{r}},Lookups!$E$3:$E${fx_last},0),"
                  f"MATCH($B$2,Lookups!$F$2:${last_fx_col}$2,0))")
    return f"Lookups!$A$3:$C${prod_last}", fx_formula


def _country_formulas(r: int, country: str, fx_formula: str) -> dict[str, str]:
    """Formulas for one month row of a country tab, as the country analysts wrote them."""
    a = f"$A{r}"
    bal = f"{DB}!$G:$G"
    crit_month = f"{DB}!$A:$A,{a}"
    crit_ctry = f"{DB}!$B:$B,$B$1"
    f = {
        "B": (f'=SUMIFS({bal},{crit_month},{crit_ctry},{DB}!$H:$H,"Personal Loans")'
              f'+SUMIFS({bal},{crit_month},{crit_ctry},{DB}!$C:$C,"CC*")'),
        "C": fx_formula.format(r=r),
        "D": f"=B{r}*C{r}",
        "E": f"=SUMIFS({DO}!$D:$D,{DO}!$A:$A,{a},{DO}!$B:$B,$B$1)",
        "F": f"=SUMIFS({DO}!$E:$E,{DO}!$A:$A,{a},{DO}!$B:$B,$B$1)",
        "G": f"=F{r}*C{r}",
        "H": f'=SUMIFS({bal},{DB}!$A:$A,EOMONTH({a},-1),{crit_ctry},{DB}!$E:$E,">=30")',
        "I": f"=IFERROR(H{r}/SUMIFS({bal},{crit_month},{crit_ctry}),0)",
        "J": f'=SUMIFS({bal},{crit_month},{crit_ctry},{DB}!$E:$E,">=90")',
        "K": f"=IFERROR(J{r}/SUMIFS({bal},{crit_month},{crit_ctry}),0)",
        "L": f"=SUMIFS({DA}!$D:$D,{DA}!$A:$A,{a},{DA}!$B:$B,$B$1)",
        "M": f"=SUMIFS({DA}!$E:$E,{DA}!$A:$A,{a},{DA}!$B:$B,$B$1)",
        "N": f"=IFERROR(M{r}/L{r},0)",
        "O": f"=SUMIFS({DB}!$F:$F,{crit_month},{crit_ctry})",
        "P": f"=IFERROR(D{r}/O{r},0)",
    }
    if country == "RO":
        f["D"] = f"=B{r}*0.2012"
        f["G"] = f"=F{r}*0.2012"
    return f


def _country_tab(ws, country: cfg.Country, months: list[date], fx_formula: str) -> None:
    ws["A1"], ws["B1"] = "Country", country.code
    ws["A2"], ws["B2"] = "Currency", country.currency
    ws["A1"].font = ws["A2"].font = Font(bold=True)
    ws["C1"] = country.name
    _header(ws, 4, [h for h, _ in COUNTRY_COLUMNS])
    ws.row_dimensions[4].height = 45
    for i, me in enumerate(months):
        r = FIRST_ROW + i
        ws.cell(row=r, column=1, value=me).number_format = "yyyy-mm-dd"
        for col, formula in _country_formulas(r, country.code, fx_formula).items():
            ws[f"{col}{r}"] = formula
        for j, (_, fmt) in enumerate(COUNTRY_COLUMNS, start=1):
            ws.cell(row=r, column=j).number_format = fmt
    for j in range(1, len(COUNTRY_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 15
    ws.freeze_panes = f"B{FIRST_ROW}"


def _summary_tab(ws, months: list[date]) -> None:
    ws["A1"] = "Group portfolio KPIs (EUR)"
    ws["A1"].font = Font(bold=True, size=14)
    headers = ["Month end", "Portfolio balance (EUR)", "New accounts", "New originations (EUR)",
               "30+ DPD rate", "90+ DPD rate", "Approval rate", "Avg balance per customer (EUR)"]
    _header(ws, 4, headers)
    ws.row_dimensions[4].height = 45
    tabs = [c.code for c in cfg.COUNTRIES]
    fmts = ["yyyy-mm-dd", FMT_MONEY, "#,##0", FMT_MONEY, FMT_RATE, FMT_RATE, FMT_RATE, FMT_MONEY]
    for i, me in enumerate(months):
        r = FIRST_ROW + i
        ws.cell(row=r, column=1, value=me)
        s = lambda col: "+".join(f"{t}!{col}{r}" for t in tabs)  # noqa: E731
        wsum = lambda col: "+".join(f"{t}!{col}{r}*{t}!D{r}" for t in tabs)  # noqa: E731
        ws[f"B{r}"] = f"={s('D')}"
        ws[f"C{r}"] = f"={s('E')}"
        ws[f"D{r}"] = f"={s('G')}"
        ws[f"E{r}"] = f"=IFERROR(({wsum('I')})/B{r},0)"
        ws[f"F{r}"] = f"=IFERROR(({wsum('K')})/B{r},0)"
        ws[f"G{r}"] = f"=IFERROR(({s('M')})/({s('L')}),0)"
        ws[f"H{r}"] = f"=IFERROR(B{r}/({s('O')}),0)"
        for j, fmt in enumerate(fmts, start=1):
            ws.cell(row=r, column=j).number_format = fmt
    for j, w in enumerate((12, 18, 12, 18, 12, 12, 12, 18), start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = f"B{FIRST_ROW}"

    last = FIRST_ROW + len(months) - 1
    chart = LineChart()
    chart.title = "Group portfolio balance (EUR)"
    chart.height, chart.width = 8, 18
    chart.add_data(Reference(ws, min_col=2, min_row=4, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=FIRST_ROW, max_row=last))
    ws.add_chart(chart, "J4")


def build_workbook(core: CoreData, path: Path) -> Path:
    months = cfg.reporting_months()
    cube = extracts.balance_cube(core)
    orig = extracts.originations(core)
    apps = extracts.applications(core)
    fx = extracts.fx_table(core)

    wb = Workbook()
    _readme(wb.active, months[-1])
    wb.active.title = "README"
    summary = wb.create_sheet("Summary")
    country_ws = {c.code: wb.create_sheet(c.code) for c in cfg.COUNTRIES}

    ws = wb.create_sheet(DB)
    _write_frame(ws, cube)
    ws.cell(row=1, column=8, value="Product line").fill = HEADER_FILL
    ws.cell(row=1, column=8).font = HEADER_FONT
    ws.column_dimensions["H"].width = 16

    _write_frame(wb.create_sheet(DO), orig)
    _write_frame(wb.create_sheet(DA), apps)
    prod_rng, fx_formula = _lookups(wb.create_sheet("Lookups"), core.products, fx)

    for r in range(2, len(cube) + 2):
        ws[f"H{r}"] = f'=IFERROR(VLOOKUP(C{r},{prod_rng},3,FALSE),"UNMAPPED")'
    for sheet in (ws, wb[DO], wb[DA]):
        for row in sheet.iter_rows(min_row=2, max_col=1):
            row[0].number_format = "yyyy-mm-dd"

    for c in cfg.COUNTRIES:
        _country_tab(country_ws[c.code], c, months, fx_formula)
    _summary_tab(summary, months)

    # No cached values are stored, so tell Excel to calculate everything when the file opens.
    wb.calculation.fullCalcOnLoad = True
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
