"""The faults found in the legacy workbook, and how each one is proved.

A fault is not "explained" by an opinion about what a formula meant. It is
explained by editing that formula in a copy of the workbook, recalculating, and
showing the published figure move by exactly the amount claimed. That is what
these fixes do: each one is a formula edit, applied to named cells, with the
number of cells it must change stated up front. If a fix stops matching the
workbook, applying it fails loudly instead of quietly doing nothing.

Fixes are applied in this order, and the order is part of the reconciliation:
where two faults touch the same figure, the interaction is attributed to the
later fix.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from portfolio_migration import config as cfg
from portfolio_migration.legacy.workbook import FIRST_ROW

COUNTRY_SHEETS = tuple(c.code for c in cfg.COUNTRIES)
MONTH_ROWS = tuple(range(FIRST_ROW, FIRST_ROW + cfg.WINDOW_MONTHS))


@dataclass(frozen=True)
class LegacyFix:
    id: str
    title: str
    short: str                   # label for the bridge walk
    sheets: tuple[str, ...]
    columns: tuple[str, ...]
    pattern: str                 # regex matched against the cell's formula
    replacement: str             # may contain {row}
    before_after: tuple[str, str]  # the formula change, as a person would write it
    expected_cells: int
    kpis: tuple[str, ...]        # every published KPI the faulty cells feed
    diagnostic_kpis: tuple[str, ...]  # the KPIs whose shape pointed at this fault
    side_effect: str             # effect beyond the obvious one, or empty
    diagnosis: str               # how the fault was located
    cause: str                   # what the formula actually does
    correct_behaviour: str       # what it should do


FIXES: tuple[LegacyFix, ...] = (
    LegacyFix(
        id="F1",
        title="A personal loan product is counted twice in the portfolio balance",
        short="double counted product",
        sheets=COUNTRY_SHEETS,
        columns=("B",),
        pattern=r'Data_Balances!\$C:\$C,"CC\*"',
        replacement='Data_Balances!$H:$H,"Credit Cards"',
        before_after=('=SUMIFS(bal,month,$A5,country,$B$1,line,"Personal Loans")'
                      '+SUMIFS(bal,month,$A5,country,$B$1,product,"CC*")',
                      '=SUMIFS(bal,month,$A5,country,$B$1,line,"Personal Loans")'
                      '+SUMIFS(bal,month,$A5,country,$B$1,line,"Credit Cards")'),
        expected_cells=72,
        kpis=("portfolio_balance_eur", "avg_balance_per_customer_eur", "dpd30_rate", "dpd90_rate"),
        diagnostic_kpis=("portfolio_balance_eur",),
        side_effect=("It also moves the group level 30+ and 90+ rates, because the Summary tab weights each "
                     "country's rate by that country's EUR balance, and those balances are overstated."),
        diagnosis=("The balance difference is zero for the first six months of the window and then grows, "
                   "in all three countries at once. Something started in that month. The product "
                   "catalogue shows Card Debt Consolidation Loan (CC_CONSOL) launched then."),
        cause=("The balance formula adds two SUMIFS: personal loans by product line, then cards by the "
               'wildcard product code "CC*". CC_CONSOL is a personal loan whose code begins with CC, so it '
               "is matched by both terms and its balance is added twice."),
        correct_behaviour=("Select cards by product line, the same way loans are selected, so the two terms "
                           "cannot overlap."),
    ),
    LegacyFix(
        id="F2",
        title="Romania is converted to EUR at a hardcoded rate",
        short="hardcoded Romania FX rate",
        sheets=("RO",),
        columns=("D", "G"),
        pattern=r"\*0\.2012",
        replacement="*C{row}",
        before_after=("=B5*0.2012", "=B5*C5"),
        expected_cells=48,
        kpis=("portfolio_balance_eur", "new_originations_eur", "avg_balance_per_customer_eur",
              "dpd30_rate", "dpd90_rate"),
        diagnostic_kpis=("portfolio_balance_eur", "new_originations_eur"),
        side_effect=("It also moves the group level rates, for the same weighting reason as F1, and it moves "
                     "the group balance even in the months before the consolidation loan launched."),
        diagnosis=("Only Romania's EUR figures differ, in every month, and the ratio between legacy and new "
                   "is constant within a month and drifts steadily across the series. A fixed factor in one "
                   "country. The RON column that displays the Treasury rate is correct, so the rate is not "
                   "the problem: the conversion is not using it."),
        cause=("The Romania tab multiplies by the literal 0.2012 instead of referencing column C. "
               "0.2012 was the RON rate when the tab was built. The displayed rate column is correct and "
               "unused, so the error is invisible on the face of the report."),
        correct_behaviour="Multiply by the month's rate from the Treasury table, as the other two tabs do.",
    ),
    LegacyFix(
        id="F3",
        title="The 30+ days past due rate uses the previous month's arrears",
        short="30+ DPD rate uses the prior month",
        sheets=COUNTRY_SHEETS,
        columns=("H",),
        pattern=r"EOMONTH\((\$A\d+),-1\)",
        replacement=r"\1",
        before_after=('=SUMIFS(bal,month,EOMONTH($A5,-1),country,$B$1,dpd,">=30")',
                      '=SUMIFS(bal,month,$A5,country,$B$1,dpd,">=30")'),
        expected_cells=72,
        kpis=("dpd30_rate",),
        diagnostic_kpis=("dpd30_rate",),
        side_effect="",
        diagnosis=("The 30+ rate is reported as 0.00% in the first month of the window, which is not "
                   "credible for a mature book, and every month's figure matches the previous month's "
                   "arrears divided by this month's balance. The 90+ rate on the same tab is fine, so it "
                   "is not the data."),
        cause=("The numerator's SUMIFS filters the balance cube on EOMONTH(month, -1), the previous month "
               "end, while the denominator uses the current month. The first month has no previous month "
               "in the cube, so the numerator is zero."),
        correct_behaviour="Filter the numerator on the same month as the denominator.",
    ),
)
FIX_BY_ID = {f.id: f for f in FIXES}


class FixDidNotApply(Exception):
    """The workbook no longer matches the fault this fix describes."""


def apply_fix(path: Path, fix: LegacyFix, out_path: Path) -> int:
    """Apply one fix to a copy of the workbook. Returns the number of cells changed."""
    wb = load_workbook(path)
    changed = 0
    for sheet in fix.sheets:
        ws = wb[sheet]
        for row in MONTH_ROWS:
            for column in fix.columns:
                cell = ws[f"{column}{row}"]
                formula = cell.value
                if not isinstance(formula, str):
                    continue
                replacement = fix.replacement.replace("{row}", str(row))
                updated = re.sub(fix.pattern, replacement, formula)
                if updated != formula:
                    cell.value = updated
                    changed += 1
    if changed != fix.expected_cells:
        raise FixDidNotApply(
            f"{fix.id} changed {changed} cells, expected {fix.expected_cells}. "
            "The workbook has changed, or the fault is not where this fix says it is.")
    wb.calculation.fullCalcOnLoad = True
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return changed


def apply_fixes_cumulatively(path: Path, fixes: tuple[LegacyFix, ...], workdir: Path) -> dict[str, Path]:
    """Build one workbook per stage: after F1, after F1 and F2, and so on.

    Returns {fix id: workbook path}. The caller recalculates each one, so the
    effect of a single fix is the difference between consecutive stages.
    """
    workdir.mkdir(parents=True, exist_ok=True)
    stages: dict[str, Path] = {}
    current = path
    for fix in fixes:
        out = workdir / f"after_{fix.id}.xlsx"
        apply_fix(current, fix, out)
        stages[fix.id] = out
        current = out
    return stages
