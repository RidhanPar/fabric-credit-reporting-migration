"""The legacy workbook must be a real formula driven workbook fed by the DWH extracts."""
import math

from openpyxl import load_workbook

from portfolio_migration import config as cfg
from portfolio_migration.legacy import extracts
from portfolio_migration.legacy.workbook import FIRST_ROW, PUBLISHED_COUNTRY_KPIS


def test_extract_totals_tie_to_the_core(core):
    cube = extracts.balance_cube(core)
    active = core.balances[core.balances.account_status == "ACTIVE"]
    assert math.isclose(cube.balance_local.sum(), active.balance_local.sum(), rel_tol=1e-9)
    assert cube.accounts.sum() == len(active)
    apps = extracts.applications(core)
    assert (apps.received == apps[["approved", "declined", "withdrawn", "incomplete"]].sum(axis=1)).all()


def test_kpi_cells_are_formulas_not_values(workbook_path):
    wb = load_workbook(workbook_path)
    for c in cfg.COUNTRIES:
        ws = wb[c.code]
        for i in range(24):
            for col in PUBLISHED_COUNTRY_KPIS:
                v = ws[f"{col}{FIRST_ROW + i}"].value
                assert isinstance(v, str) and v.startswith("="), (c.code, col, i)
    assert {"README", "Summary", "Data_Balances", "Data_Originations", "Data_Applications",
            "Lookups"} <= set(wb.sheetnames)


def test_every_published_kpi_evaluates_to_a_number(legacy_kpis):
    assert len(legacy_kpis) == 24 * 4 * 7
    assert legacy_kpis.value.map(math.isfinite).all()
    rates = legacy_kpis[legacy_kpis.kpi.str.endswith("rate")]
    assert rates.value.between(0, 1).all()


def test_new_accounts_tie_to_originations_extract(core, legacy_kpis):
    orig = extracts.originations(core).groupby(["month_end", "country_code"]).new_accounts.sum()
    got = legacy_kpis[(legacy_kpis.kpi == "new_accounts") & (legacy_kpis.scope != "GROUP")]
    for rec in got.itertuples():
        assert rec.value == orig.get((rec.month_end, rec.scope), 0)


def test_committed_legacy_figures_reproduce(legacy_kpis):
    """Every legacy number quoted in the docs comes from data/legacy/legacy_published_kpis.csv.
    Regenerating from the seed must give the same figures."""
    from pathlib import Path

    import pandas as pd

    committed = Path(__file__).parents[1] / "data" / "legacy" / "legacy_published_kpis.csv"
    old = pd.read_csv(committed, parse_dates=["month_end"])
    new = legacy_kpis.assign(month_end=pd.to_datetime(legacy_kpis.month_end))
    pd.testing.assert_frame_equal(old, new, check_exact=False, rtol=1e-9)
